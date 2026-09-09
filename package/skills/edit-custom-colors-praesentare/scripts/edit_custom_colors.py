#!/usr/bin/env python3
import argparse,json,os,re,tempfile,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
A='http://schemas.openxmlformats.org/drawingml/2006/main';P='http://schemas.openxmlformats.org/presentationml/2006/main';ET.register_namespace('a',A)
PREFIX='PCL_CustomColor_'; EMPTY_RGB='FFFFFF'; HEADER='Skill edit-custom-colors-praesentare by Peter Claus Lamprecht, https://praesentare.com'

def parse_theme(data):
 root=ET.fromstring(data);scheme={};cs=root.find(f'.//{{{A}}}clrScheme')
 if cs is not None:
  for h in list(cs):
   if not list(h):continue
   n=list(h)[0];k=n.tag.rsplit('}',1)[-1];v=n.get('val') if k=='srgbClr' else n.get('lastClr') if k=='sysClr' else None
   if v:scheme[h.tag.rsplit('}',1)[-1]]=v.upper()
 out=[];cl=root.find(f'{{{A}}}custClrLst')
 if cl is not None:
  for pos,x in enumerate(cl.findall(f'{{{A}}}custClr'),1):
   name=x.get('name','');n=list(x)[0] if list(x) else None;rgb=None
   if n is not None:
    k=n.tag.rsplit('}',1)[-1];rgb=n.get('val') if k=='srgbClr' else n.get('lastClr') if k=='sysClr' else scheme.get(n.get('val'))
   rgb=rgb.upper() if rgb else None
   out.append({'slot':pos,'name':name,'rgb':rgb,'empty':name=='' and rgb==EMPTY_RGB})
 return root,scheme,out

def read_theme(path):
 with zipfile.ZipFile(path) as z:return parse_theme(z.read('ppt/theme/theme1.xml'))

def choose_layout(prs):
 for x in prs.slide_layouts:
  if (x.name or '').strip().lower() in {'blank','leer'}:return x
 return min(prs.slide_layouts,key=lambda x:len(x.placeholders))

def remove_placeholders(slide):
 for shape in list(slide.shapes):
  if shape.is_placeholder:
   el=shape._element;el.getparent().remove(el)

def load_grid(path,colors):
 from pptx import Presentation
 from pptx.util import Inches,Pt
 from pptx.enum.shapes import MSO_SHAPE
 from pptx.dml.color import RGBColor
 prs=Presentation(str(path));slide=prs.slides.add_slide(choose_layout(prs));remove_placeholders(slide)
 margin=Inches(.5);gap=Inches(.08);header_h=Inches(.35);header_gap=Inches(.18)
 box=slide.shapes.add_textbox(margin,Inches(.3),prs.slide_width-2*margin,header_h);tf=box.text_frame;tf.clear();tf.paragraphs[0].text=HEADER
 # No explicit font face: PowerPoint applies the presentation/theme default font.
 max_w=prs.slide_width-2*margin;top=Inches(.3)+header_h+header_gap;max_h=prs.slide_height-top-Inches(.45)
 side=min((max_w-gap*9)//10,(max_h-gap*4)//5);total_w=side*10+gap*9;total_h=side*5+gap*4;left=(prs.slide_width-total_w)//2
 for i in range(50):
  s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left+i%10*(side+gap),top+i//10*(side+gap),side,side);s.name=f'{PREFIX}{i+1:02d}'
  c=colors[i] if i<len(colors) else {'empty':True,'rgb':None}
  if c.get('rgb') and not c.get('empty'):s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(c['rgb'])
  else:s.fill.background()
  s.line.color.rgb=RGBColor(0,0,0);s.line.width=Pt(.75)
 prs.save(str(path));return len(prs.slides)

def shape_rgb(sp,scheme):
 solid=sp.find(f'{{{P}}}spPr/{{{A}}}solidFill')
 if solid is None or not list(solid):return None
 n=list(solid)[0];k=n.tag.rsplit('}',1)[-1];v=n.get('val') if k=='srgbClr' else n.get('lastClr') if k=='sysClr' else scheme.get(n.get('val'))
 return v.upper() if v and re.fullmatch(r'[0-9A-Fa-f]{6}',v) else None

def shape_text(sp):
 return ''.join((t.text or '') for t in sp.findall(f'.//{{{A}}}t')).strip()

def read_grid(path,scheme):
 with zipfile.ZipFile(path) as z:
  slides=sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda n:int(re.search(r'slide(\d+)',n).group(1)),reverse=True)
  for member in slides:
   root=ET.fromstring(z.read(member));found={}
   for sp in root.findall(f'.//{{{P}}}sp'):
    n=sp.find(f'{{{P}}}nvSpPr/{{{P}}}cNvPr');name=n.get('name','') if n is not None else ''
    if re.fullmatch(PREFIX+r'\d{2}',name):found[name]={'rgb':shape_rgb(sp,scheme),'text':shape_text(sp)}
   if len(found)==50:return [{'slot':i,**found[f'{PREFIX}{i:02d}']} for i in range(1,51)]
 raise RuntimeError('Keine vollständige Farbmatrix gefunden.')

def write_theme(path,slots):
 occupied=[x['slot'] for x in slots if x['rgb']];last=max(occupied,default=0);to_write=slots[:last]
 fd,tmp=tempfile.mkstemp(suffix='.pptx',dir=path.parent);os.close(fd);tmp=Path(tmp)
 try:
  with zipfile.ZipFile(path) as zi,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zo:
   for item in zi.infolist():
    data=zi.read(item.filename)
    if item.filename=='ppt/theme/theme1.xml':
     root=ET.fromstring(data);old=root.find(f'{{{A}}}custClrLst')
     if old is not None:root.remove(old)
     if to_write:
      cl=ET.Element(f'{{{A}}}custClrLst')
      for x in to_write:
       rgb=x['rgb'] or EMPTY_RGB;name=(x['text'] if x['rgb'] and x['text'] else f'#{rgb}' if x['rgb'] else '')
       cc=ET.SubElement(cl,f'{{{A}}}custClr',{'name':name});ET.SubElement(cc,f'{{{A}}}srgbClr',{'val':rgb})
      ext=root.find(f'{{{A}}}extLst');root.insert(list(root).index(ext) if ext is not None else len(root),cl)
     data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    zo.writestr(item,data)
  os.replace(tmp,path)
 finally:
  if tmp.exists():tmp.unlink()
 return len(to_write),last

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--presentation',required=True);ap.add_argument('--mode',choices=['load','laden','save','speichern'],default='load',nargs='?');a=ap.parse_args();mode='save' if a.mode in ('save','speichern') else 'load';p=Path(a.presentation).resolve();r={'mode':mode,'success':False}
 try:
  _,scheme,colors=read_theme(p)
  if mode=='load':r.update(success=True,custom_color_entries_read=len(colors),occupied_colors=sum(not x['empty'] for x in colors),slide_count_after=load_grid(p,colors))
  else:
   slots=read_grid(p,scheme);count,last=write_theme(p,slots);r.update(success=True,occupied_colors=sum(bool(x['rgb']) for x in slots),custom_color_entries_written=count,last_defined_slot=last)
 except Exception as e:r['error']=f'{type(e).__name__}: {e}'
 print(json.dumps(r,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
